"""ワンクリック実行テスト用の合成データ (実データは含まない)。

- 出張精算CSV: 楽楽精算の実際の出力と同じ32列のヘッダ + 架空の明細 (cp932)
- 出勤簿xlsx: 楽楽勤怠の「出勤簿（日別詳細）」と同じ形 (1行目タイトル・2行目空・3行目ヘッダ、
  時刻は 'HH:MM' 文字列、日付は datetime、シート 一般 / 管理職)
"""
from __future__ import annotations

import calendar
import datetime as dt
import io

from openpyxl import Workbook

EXPENSE_HEADER = (
    "ヘッダ情報:伝票No.(伝票No.),ヘッダ情報:出張申請伝票No.(申請No.),ヘッダ情報:行数,明細情報:明細No.,"
    "明細情報:明細日付(日付),明細情報:明細時刻(開始)(時刻),明細情報:明細時刻(終了)(時刻),明細情報:出発地(出発),"
    "明細情報:到着地(到着),明細情報:交通機関(交通機関他),明細情報:金額(金額),明細情報:証票(領収証),"
    "明細情報:手当1CD(日当(選択)),明細情報:手当2CD(宿泊料(選択)),明細情報:手当3CD(滞在費補助(選択)),"
    "明細情報:フリー１(備考),明細情報:勘定科目名,明細情報:小計(小計),ヘッダ情報:手当計(手当計),ヘッダ情報:合計(合計),"
    "ヘッダ情報:申請者CD(申請者),ヘッダ情報:申請者名(申請者),ヘッダ情報:承認実行者1名,ヘッダ情報:承認日1,"
    "ヘッダ情報:承認実行者2名,ヘッダ情報:承認日2,ヘッダ情報:承認実行者3名,ヘッダ情報:承認日3,"
    "ヘッダ情報:承認実行者4名,ヘッダ情報:承認日4,ヘッダ情報:承認実行者5名,ヘッダ情報:承認日5"
)
EMP_ID = "AH900001"
EMP_NAME = "試験　一郎"


def expense_csv_bytes(dates=("2026/08/28", "2026/09/03")):
    """架空の出張1件 (往路・復路の2明細)。dates は各明細の日付。"""
    rows = [EXPENSE_HEADER]
    legs = [("東京", "大阪", "09:00", "11:30"), ("大阪", "東京", "17:00", "19:30")]
    for i, day in enumerate(dates, 1):
        dep, arr, t1, t2 = legs[(i - 1) % 2]
        rows.append(",".join([
            "90001", "80001", str(len(dates)), str(i), day, t1, t2, dep, arr, "JR", "13500", "無",
            "", "", "", "テスト出張", "旅費交通費", "13500", "0", str(13500 * len(dates)),
            EMP_ID, EMP_NAME, "", "", "", "", "", "", "", "", "", "",
        ]))
    return ("\r\n".join(rows) + "\r\n").encode("cp932")


ATTENDANCE_HEADER = ["社員番号", "氏名", "部門", "役職", "日付", "曜日", "カレンダー", "出勤時刻", "退勤時刻",
                     "テレワーク出勤時刻", "テレワーク退勤時刻", "移動開始", "移動終了", "移動2開始", "移動2終了",
                     "申請内容", "休出時間", "備考"]
_WD = "月火水木金土日"


def attendance_xlsx_bytes(ym, emp_id=EMP_ID, name=EMP_NAME, late_out=None):
    """楽楽勤怠の出勤簿（日別詳細）と同じ形の1か月分。late_out={日: '24:30'} で深夜退勤を入れる。"""
    y, m = [int(x) for x in ym.split("-")]
    wb = Workbook()
    for idx, sheet in enumerate(("一般", "管理職")):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = sheet
        ws.append(["出勤簿（日別詳細） %d年%02d月" % (y, m)])
        ws.append([])
        ws.append(ATTENDANCE_HEADER)
        if sheet != "一般":
            continue
        for d in range(1, calendar.monthrange(y, m)[1] + 1):
            day = dt.datetime(y, m, d)
            weekend = day.weekday() >= 5
            out = (late_out or {}).get(d, "18:00")
            ws.append([emp_id, name, "技術部", "社員", day, _WD[day.weekday()], "休日" if weekend else "平日",
                       None if weekend else "09:00", None if weekend else out, None, None,
                       None, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def engine_format_xlsx_bytes():
    """エンジンがそのまま読める形 (1行目ヘッダ・時刻は datetime)。正規化で変わらないことの確認用。"""
    wb = Workbook()
    ws = wb.active
    ws.append(ATTENDANCE_HEADER)
    ws.append([EMP_ID, EMP_NAME, "技術部", "社員", dt.datetime(2026, 9, 1), "火", "平日",
               dt.datetime(2026, 9, 1, 9, 0), dt.datetime(2026, 9, 1, 18, 0)] + [None] * 9)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
