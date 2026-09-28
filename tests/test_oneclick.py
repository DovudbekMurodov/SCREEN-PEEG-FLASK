"""ワンクリック実行: 対象月の算出・出勤簿の正規化・実エンジンでのチェックシート作成 (ブラウザ不要)。"""
from __future__ import annotations

import base64
import html
import contextlib
import datetime as dt
import io
import json
import os
import re
import sys

import pytest

import app as appmod
import rakuraku_web
from rakuraku.oneclick import months_from_expense_csv, normalize_attendance_xlsx, plan_kintai_months
from tests.oneclick_fixtures import (
    EMP_NAME,
    attendance_xlsx_bytes,
    engine_format_xlsx_bytes,
    expense_csv_bytes,
)

if appmod.SRC_DIR not in sys.path:
    sys.path.insert(0, appmod.SRC_DIR)
from loaders.attendance_loader import load_attendance  # noqa: E402  (src エンジンの実ローダ)


def _load(tmp_path, data, name="a.xlsx"):
    path = tmp_path / name
    path.write_bytes(data)
    return load_attendance([str(path)])


def sse_events(body):
    out = []
    for chunk in body.split("\n\n"):
        ev, data = None, ""
        for line in chunk.split("\n"):
            if line.startswith("event:"):
                ev = line[6:].strip()
            elif line.startswith("data:"):
                data += line[5:].strip()
        if ev:
            out.append((ev, json.loads(data) if data else None))
    return out


# ------------------------------------------------------------------ 対象月
def test_months_cover_detail_date_range_newest_first():
    assert months_from_expense_csv(expense_csv_bytes(("2026/07/30", "2026/09/02"))) == ["2026-09", "2026-08", "2026-07"]


def test_months_single_month_and_year_boundary():
    assert months_from_expense_csv(expense_csv_bytes(("2026/09/03", "2026/09/04"))) == ["2026-09"]
    assert months_from_expense_csv(expense_csv_bytes(("2025/12/30", "2026/01/05"))) == ["2026-01", "2025-12"]


def test_months_utf8_bom_and_missing_column():
    text = expense_csv_bytes().decode("cp932")
    assert months_from_expense_csv(text.encode("utf-8-sig")) == ["2026-09", "2026-08"]
    assert months_from_expense_csv("伝票No,金額\r\n1,100\r\n".encode("cp932")) == []


BASE = ["2026-09", "2026-08"]  # 9月に実行したときの 当月・前月


def test_plan_always_includes_current_and_previous_month():
    # お客様の要望: CSVが9月分だけでも、8月・9月の2か月を取得する
    assert plan_kintai_months(["2026-09"], BASE, 4) == (["2026-09", "2026-08"], [])
    assert plan_kintai_months([], BASE, 4) == (["2026-09", "2026-08"], [])


def test_plan_adds_older_csv_months_and_caps():
    assert plan_kintai_months(["2026-09", "2026-08", "2026-07"], BASE, 4) == (["2026-09", "2026-08", "2026-07"], [])
    csv = ["2026-09", "2026-08", "2026-07", "2026-06", "2026-05"]
    assert plan_kintai_months(csv, BASE, 4) == (["2026-09", "2026-08", "2026-07", "2026-06"], ["2026-05"])
    assert plan_kintai_months(csv, BASE, 1) == (["2026-09", "2026-08"], ["2026-07", "2026-06", "2026-05"])  # 当月・前月は上限でも削らない


# ------------------------------------------------------------------ 出勤簿の正規化
def test_raw_kintai_export_is_unreadable_by_engine_and_normalized_is_readable(tmp_path):
    raw = attendance_xlsx_bytes("2026-09")
    assert _load(tmp_path, raw) == []  # 不具合の固定: タイトル行のせいで1行も読めない
    days = _load(tmp_path, normalize_attendance_xlsx(raw), "b.xlsx")
    assert len(days) == 30 and {d.name_raw for d in days} == {EMP_NAME}
    worked = [d for d in days if d.clock_in]
    assert worked and worked[0].clock_in == dt.datetime(2026, 9, worked[0].work_date.day, 9, 0)
    assert all(d.clock_out.hour == 18 for d in worked)


def test_after_midnight_clock_out_rolls_to_next_day(tmp_path):
    days = _load(tmp_path, normalize_attendance_xlsx(attendance_xlsx_bytes("2026-09", late_out={1: "24:30"})))
    first = next(d for d in days if d.work_date == dt.date(2026, 9, 1))
    assert first.clock_out == dt.datetime(2026, 9, 2, 0, 30)


def test_time_objects_are_converted(tmp_path):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(attendance_xlsx_bytes("2026-09")))
    wb["一般"].cell(row=4, column=8).value = dt.time(8, 45)  # 1日の 出勤時刻 を time 値に
    buf = io.BytesIO()
    wb.save(buf)
    days = _load(tmp_path, normalize_attendance_xlsx(buf.getvalue()))
    first = next(d for d in days if d.work_date == dt.date(2026, 9, 1))
    assert first.clock_in == dt.datetime(2026, 9, 1, 8, 45)


def test_engine_format_file_is_returned_unchanged():
    data = engine_format_xlsx_bytes()
    assert normalize_attendance_xlsx(data) is data
    fixed = normalize_attendance_xlsx(attendance_xlsx_bytes("2026-09"))
    assert normalize_attendance_xlsx(fixed) is fixed  # 2回目は何もしない


# ------------------------------------------------------------------ /run (手動アップロード) の回帰
def _attendance_count(log):
    lines = log.splitlines()
    for i, line in enumerate(lines):
        if "勤怠" in line and "読込" in line:
            m = re.search(r"->\s*(\d+)\s*件", lines[i + 1])
            return int(m.group(1)) if m else None
    return None


def test_run_upload_of_raw_kintai_export_now_loads_attendance(client):
    r = client.post("/run", data={
        "expense": (io.BytesIO(expense_csv_bytes()), "出張精算_x.csv"),
        "attendance": (io.BytesIO(attendance_xlsx_bytes("2026-09")), "出勤簿_日別詳細_x.xlsx"),
        "approver": "",
    }, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)[-800:]
    body = html.unescape(r.get_data(as_text=True))  # ログは <pre> 内で -&gt; にエスケープされる
    assert _attendance_count(body) == 30  # 修正前は 0 件 (全件 未確認)


# ------------------------------------------------------------------ /oneclick/stream + 実エンジン
@contextlib.contextmanager
def _fake_session(**kwargs):
    yield object()


class _FakeSeisan:
    STEPS = rakuraku_web.SeisanClient.STEPS
    calls = []
    dates = ("2026/08/28", "2026/09/03")

    def __init__(self, *a, **k):
        pass

    def download(self, params, on_step=None):
        _FakeSeisan.calls.append(params)
        for key, _label in self.STEPS:
            on_step(key, "running")
            on_step(key, "done")
        return "出張精算_20260928_000000.csv", expense_csv_bytes(_FakeSeisan.dates)

    def logout(self):
        pass


class _FakeKintai:
    steps_for = staticmethod(rakuraku_web.KintaiClient.steps_for)
    empty = ()
    months_seen = []

    def __init__(self, *a, **k):
        self.skipped = []

    def download(self, params, progress=None, on_step=None):
        _FakeKintai.months_seen.append(list(params.months))
        files = []
        for ym in params.months:
            if ym in _FakeKintai.empty:
                self.skipped.append(ym)
                on_step("export:%s" % ym, "skipped", "%s の出勤簿データがありません（この月はスキップしました）。" % ym)
                continue
            files.append(("%s_出勤簿_日別詳細_20260928.xlsx" % ym, attendance_xlsx_bytes(ym)))
            on_step("export:%s" % ym, "done")
        return files

    def logout(self):
        pass


@pytest.fixture
def fake_clients(monkeypatch):
    monkeypatch.setattr(rakuraku_web, "playwright_available", lambda: True)
    monkeypatch.setattr(rakuraku_web, "browser_session", _fake_session)
    monkeypatch.setattr(rakuraku_web, "SeisanClient", _FakeSeisan)
    monkeypatch.setattr(rakuraku_web, "KintaiClient", _FakeKintai)
    monkeypatch.setattr(rakuraku_web, "today_jst", lambda: dt.date(2026, 9, 28))  # 実行日を固定
    _FakeSeisan.calls[:] = []
    _FakeSeisan.dates = ("2026/08/28", "2026/09/03")
    _FakeKintai.months_seen[:] = []
    _FakeKintai.empty = ()


FORM = {"s_login_id": "AZ1", "s_password": "sp", "k_company_code": "PEEG", "k_login_id": "AZ1", "k_password": "kp"}


def _oneclick_jobs():
    return [d for d in os.listdir(appmod.JOBS_DIR) if d.startswith("oc")] if os.path.isdir(appmod.JOBS_DIR) else []


def test_oneclick_runs_seisan_kintai_and_real_engine(client, fake_clients):
    before = set(_oneclick_jobs())
    events = sse_events(client.post("/oneclick/stream", data=FORM).get_data(as_text=True))
    kinds = [e for e, _ in events]
    assert "error" not in kinds, events
    # 対象月は CSV の明細日付 (8/28〜9/3) から自動で決まり、チェック前に挿入される
    assert _FakeKintai.months_seen == [["2026-09", "2026-08"]]
    inserted = [d for e, d in events if e == "steps" and d.get("before") == "check"][0]
    assert [s["key"] for s in inserted["steps"]] == ["k:export:2026-09", "k:export:2026-08"]
    assert {"key": "check", "status": "done", "detail": None} in [d for e, d in events if e == "step"]

    result = [d for e, d in events if e == "result"][0]
    names = [f["filename"] for f in result["files"]]
    assert any(n.endswith(".xlsx") for n in names) and any(n.endswith("_ja.html") for n in names)
    assert result["summary"] and all("count" in s for s in result["summary"])
    assert base64.b64decode(result["files"][0]["b64"])[:2] == b"PK" or names[0].endswith(".html")
    assert [f["filename"] for f in result["sources"]][0].endswith(".csv")
    assert _attendance_count(result["log"]) == 61  # 9月30日 + 8月31日 が読めている
    assert result["skipped"] == [] and result["dropped"] == []
    assert set(_oneclick_jobs()) == before  # 作業フォルダはサーバに残さない


def test_oneclick_fetches_previous_month_even_if_csv_has_only_current_month(client, fake_clients):
    _FakeSeisan.dates = ("2026/09/10", "2026/09/11")
    events = sse_events(client.post("/oneclick/stream", data=FORM).get_data(as_text=True))
    assert _FakeKintai.months_seen == [["2026-09", "2026-08"]]
    assert _attendance_count([d for e, d in events if e == "result"][0]["log"]) == 61


def test_oneclick_adds_older_month_from_csv(client, fake_clients):
    _FakeSeisan.dates = ("2026/07/30", "2026/09/02")
    sse_events(client.post("/oneclick/stream", data=FORM).get_data(as_text=True))
    assert _FakeKintai.months_seen == [["2026-09", "2026-08", "2026-07"]]


def test_oneclick_skipped_month_still_creates_check_sheet(client, fake_clients):
    _FakeKintai.empty = ("2026-08",)
    events = sse_events(client.post("/oneclick/stream", data=FORM).get_data(as_text=True))
    result = [d for e, d in events if e == "result"][0]
    assert result["skipped"] == ["2026-08"]
    assert _attendance_count(result["log"]) == 30


def test_oneclick_param_error_and_unknown_approver(client, fake_clients):
    events = sse_events(client.post("/oneclick/stream", data={"s_login_id": "AZ1"}).get_data(as_text=True))
    assert events[0][0] == "error"
    form = dict(FORM, approver="存在しない 名前")
    events = sse_events(client.post("/oneclick/stream", data=form).get_data(as_text=True))
    assert [d for e, d in events if e == "result"][0]["approver"] == "全員（絞り込みなし）"


def test_oneclick_page_renders_with_shared_credential_keys(client):
    body = client.get("/oneclick").get_data(as_text=True)
    assert 'data-rr-ns="seisan" data-rr-store="login_id"' in body
    assert 'data-rr-ns="kintai" data-rr-store="password"' in body
    assert 'href="/oneclick"' in client.get("/").get_data(as_text=True)  # 一番目のメニュー + 案内
