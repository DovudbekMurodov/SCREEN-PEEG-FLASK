"""ワンクリック実行: 楽楽精算CSV → 対象月 → 楽楽勤怠の出勤簿 → チェックシート (src エンジン)。

src/ のエンジンは変更しない。楽楽勤怠の「出勤簿（日別詳細）」はそのままだと
エンジンが1行も読めない (1行目がタイトルでヘッダは3行目、時刻が 'HH:MM' の文字列)
ため、ここでエンジンが読める形に整えてから渡す。
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import os
import re
import shutil
import subprocess
import time
import unicodedata
import uuid

import openpyxl

from rakuraku.errors import CheckSheetFailed

# エンジン (attendance_loader) が parse_excel_dt で読む列。datetime か
# 'YYYY-MM-DD HH:MM' 文字列しか受け付けない。
TIME_COLUMNS = ("出勤時刻", "退勤時刻", "テレワーク出勤時刻", "テレワーク退勤時刻")
HEADER_SCAN_ROWS = 30
_HHMM = re.compile(r"^\s*(\d{1,2}):(\d{2})(?::(\d{2}))?\s*$")
_YMD = re.compile(r"(\d{4})\s*[/\-.年]\s*(\d{1,2})\s*[/\-.月]\s*(\d{1,2})")


# --------------------------------------------------------------------- 出張精算CSV
def _decode_csv(data):
    # type: (bytes) -> str
    for enc in ("cp932", "utf-8-sig"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("cp932", errors="replace")


def _to_date(value):
    # type: (object) -> dt.date | None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        m = _YMD.search(unicodedata.normalize("NFKC", value))
        if m:
            try:
                return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                return None
    return None


def expense_detail_dates(data):
    # type: (bytes) -> list
    """出張精算CSV の「明細日付」列の日付 (読めない行は無視)。列が無ければ []。"""
    reader = csv.reader(io.StringIO(_decode_csv(data)))
    header = next(reader, None) or []
    col = next((i for i, h in enumerate(header) if "明細日付" in (h or "")), None)
    if col is None:
        return []
    days = []
    for row in reader:
        if col < len(row):
            d = _to_date(row[col])
            if d:
                days.append(d)
    return days


def months_from_expense_csv(data):
    # type: (bytes) -> list
    """明細日付 の最小〜最大を含む YYYY-MM を新しい順で返す (出勤簿が必要な月)。"""
    days = expense_detail_dates(data)
    if not days:
        return []
    lo, hi = min(days), max(days)
    months = []
    y, m = lo.year, lo.month
    while (y, m) <= (hi.year, hi.month):
        months.append("%04d-%02d" % (y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return sorted(months, reverse=True)


# --------------------------------------------------------------------- 出勤簿 xlsx
def _clock(day, value):
    # type: (dt.date, object) -> dt.datetime | None
    """時刻セルを 日付+時刻 の datetime に。変換不要・不能なら None (元の値のまま)。"""
    if value is None or value == "" or isinstance(value, dt.datetime):
        return None
    if isinstance(value, dt.time):
        return dt.datetime.combine(day, value)
    if isinstance(value, str):
        m = _HHMM.match(unicodedata.normalize("NFKC", value))
        if not m:
            return None
        h, mi, se = int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)
        if h >= 48 or mi >= 60 or se >= 60:
            return None
        extra, h = divmod(h, 24)  # '24:30' 等の翌日表記
        return dt.datetime.combine(day, dt.time(h, mi, se)) + dt.timedelta(days=extra)
    return None


def normalize_attendance_xlsx(data):
    # type: (bytes) -> bytes
    """楽楽勤怠の出勤簿をエンジンが読める形に整える。

    各シートで「社員番号」「氏名」を含む行をヘッダとみなしてその上の行を削除し、
    時刻列の 'HH:MM' 文字列・time 値を 日付+時刻 に変換する。
    既にエンジン形式 (1行目がヘッダ、時刻が日時) のファイルは元のバイト列をそのまま返す。
    """
    wb = openpyxl.load_workbook(io.BytesIO(data))
    changed = False
    for ws in wb.worksheets:
        header_row = None
        for r, row in enumerate(ws.iter_rows(max_row=min(ws.max_row, HEADER_SCAN_ROWS), values_only=True), 1):
            names = {str(v).strip() for v in row if v is not None}
            if "社員番号" in names and "氏名" in names:
                header_row = r
                break
        if header_row is None:
            continue
        if header_row > 1:
            ws.delete_rows(1, header_row - 1)
            changed = True
        cols = {}
        for i, c in enumerate(ws[1]):
            if c.value is not None:
                cols.setdefault(str(c.value).strip(), i)  # エンジンと同じく最初の列を採用
        if "日付" not in cols:
            continue
        date_col = cols["日付"]
        time_cols = [cols[n] for n in TIME_COLUMNS if n in cols]
        for row in ws.iter_rows(min_row=2):
            day = _to_date(row[date_col].value) if date_col < len(row) else None
            if day is None:
                continue
            for ci in time_cols:
                if ci >= len(row):
                    continue
                new = _clock(day, row[ci].value)
                if new is not None:
                    row[ci].value = new
                    row[ci].number_format = "yyyy-mm-dd hh:mm"
                    changed = True
    if not changed:
        return data
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def normalize_attendance_file(path):
    # type: (str) -> bool
    """ファイルをその場で正規化する (/run のアップロード用)。変更したら True。"""
    with open(path, "rb") as fh:
        data = fh.read()
    fixed = normalize_attendance_xlsx(data)
    if fixed is data:
        return False
    with open(path, "wb") as fh:
        fh.write(fixed)
    return True


# --------------------------------------------------------------------- エンジン実行
_MIME = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".html": "text/html",
    ".csv": "text/csv",
}


def mime_of(name):
    # type: (str) -> str
    return _MIME.get(os.path.splitext(name)[1].lower(), "application/octet-stream")


def run_engine(engine, csv_bytes, attendance, approver=""):
    # type: (object, bytes, list, str) -> dict
    """src エンジンでチェックシートを作る。

    engine: app.py が app.config["CHECKSHEET_ENGINE"] に渡すヘルパ群。
    attendance: [(YYYY-MM, 正規化済み xlsx bytes)]。
    戻り値 {"summary": [(判定, 件数)], "files": [(名前, bytes)], "log": str}。
    作業フォルダは成功・失敗にかかわらず必ず削除する (サーバに何も残さない)。
    """
    work = engine.prepare("oc" + uuid.uuid4().hex[:10])
    try:
        stamp = time.strftime("%Y%m%d_%H%M%S")
        with open(os.path.join(work, "expenses", "出張精算_%s.csv" % stamp), "wb") as fh:
            fh.write(csv_bytes)
        paths = []
        for ym, data in attendance:
            path = os.path.join(work, "attendance", "出勤簿_%s.xlsx" % ym)
            with open(path, "wb") as fh:
                fh.write(data)
            paths.append(path)
        # エンジンは既定で最新の出勤簿1つしか読まないため、全月を明示する。
        cfg = os.path.join(work, "oneclick_config.json")
        with open(cfg, "w", encoding="utf-8") as fh:
            json.dump({"attendance_paths": paths}, fh, ensure_ascii=False)
        cmd = [engine.python(), os.path.join(engine.src_dir, "main.py"),
               "--no-pause", "--approver", approver, "--config", cfg]
        env = dict(os.environ, CHECKSHEET_ROOT=work, PYTHONIOENCODING="utf-8")
        try:
            # stdin を閉じる: マスタのパスワード入力待ち (getpass) で固まらないように。
            proc = subprocess.run(cmd, capture_output=True, text=True, env=env,
                                  stdin=subprocess.DEVNULL, timeout=engine.timeout)
        except subprocess.TimeoutExpired:
            raise CheckSheetFailed("engine timeout", user_message="チェックシートの作成が時間内に終わりませんでした。")
        log = engine.clean_log((proc.stdout or "") + (proc.stderr or ""), work)
        if proc.returncode != 0:
            raise CheckSheetFailed("engine exit %s" % proc.returncode, log=log)
        stamp_dir, names = engine.results(work)
        if not names:
            raise CheckSheetFailed("no output", log=log)
        files = []
        for name in names:
            with open(os.path.join(work, "out", stamp_dir, name), "rb") as fh:
                files.append((name, fh.read()))
        return {"summary": engine.summary(log), "files": files, "log": log}
    finally:
        shutil.rmtree(work, ignore_errors=True)
